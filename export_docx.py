"""Export the filled manuscript to Word, with figures and tables placed inline.

  python export_docx.py                # real-data outputs
  python export_docx.py --synthetic

Produces manuscript/manuscript_<tag>.docx. Co-authors can then track-change it,
which is why Word rather than PDF.

Figures are inserted immediately after the subsection that first refers to them,
and result tables are generated from the CSVs in outputs/tables so the document
can never drift from the numbers the pipeline produced.
"""
from __future__ import annotations

import argparse
import re
import shutil
import subprocess
import sys
from pathlib import Path

import pandas as pd

sys.path.insert(0, str(Path(__file__).resolve().parent / "src"))
from cdh import config as C  # noqa: E402

# Each entry anchors to the heading that FOLLOWS it, so the
# figure lands at the end of the section that refers to it. Numbers are assigned
# from anchor order, shared by callouts, captions and the submission packager.
# About half of these are the paper's argument (plus the raw climate maps
# themselves, kept in the body rather than supplementary since a compound
# drought-heat paper should show the continuous rainfall/temperature fields
# it is built on, not just the derived indices); the rest are the evidence
# behind it. `main` marks the figures that belong in the manuscript body;
# everything else is real, checked work that a reviewer can still see, just
# in the Supplementary Material instead of competing with the headline
# result for the reader's attention.
FIGURES = [
    ("### 2.2 Food insecurity outcome", "map0_locator",
     "Location of the study area. (a) Zimbabwe within Africa. (b) Regional context. "
     "(c) Provinces and district/council boundaries, shown in an Albers equal-area "
     "projection. Administrative boundaries: geoBoundaries gbOpen (CC BY 4.0); "
     "continental and regional context: Natural Earth 1:110 million, public domain.",
     True),
    ("## 3. Methods", "fig_study_design",
     "Season-year timing. The growing season (October of Y-1 to March of Y) is "
     "when exposure (SPEI-6, growing-season hot days) is measured; the outcome "
     "(mean lean-season IPC phase) is measured in the lean season of Y+1, nine "
     "to fourteen months later, a lag fixed by the agricultural calendar.",
     False),
    ("### 2.3 Exposure", "fig0_national_trend",
     "National record, 2009-2025. (a) Mean lean-season IPC phase, population-"
     "weighted, with the interquartile range across districts shaded. "
     "(b) Share of districts classified compound drought-heat each season.",
     True),
    ("### 2.2 Food insecurity outcome", "map1_study_area",
     "Zimbabwe climate gradients. (a) Mean growing-season rainfall by district. "
     "(b) Mean growing-season daily maximum temperature. Province boundaries in grey. "
     "Albers equal-area projection. Boundaries: geoBoundaries gbOpen (CC BY 4.0).",
     True),
    ("### 2.4 Temporal alignment", "map4_trends",
     "Descriptive climate trend, not the paper's estimand: the two-way fixed "
     "effects model absorbs any trend common across districts. (a) Growing-season "
     "rainfall trend. (b) Growing-season maximum temperature trend, both per decade "
     "over the climate record.",
     False),
    ("### 2.4 Temporal alignment", "map_rainfall_periods",
     "Growing-season rainfall by period, the raw continuous exposure variable, "
     "prior to any standardisation or drought classification. Native 5.6 km CHIRPS "
     "v3 grid, not a district choropleth; boundaries shown for reference only. "
     "Small multiples on one shared colour scale.",
     True),
    ("### 2.4 Temporal alignment", "map_temperature_periods",
     "Growing-season mean daily maximum temperature by period, the raw continuous "
     "exposure variable, prior to any hot-day thresholding. Native 11.1 km "
     "ERA5-Land grid, not a district choropleth; boundaries shown for reference "
     "only. Same period bins and colour-scale convention as the rainfall-by-period "
     "figure.",
     True),
    ("### 3.3 Compound specification", "map_moisture_deficit_periods",
     "Moisture deficit (SPEI-6, reversed, z) by period: the continuous, district-"
     "standardised index itself, before the binary drought threshold below "
     "is applied to it.",
     False),
    ("### 3.3 Compound specification", "map_heat_periods",
     "Growing-season hot days (z) by period: the continuous, district-standardised "
     "heat exposure itself, before the binary hot-season threshold below "
     "is applied to it.",
     False),
    ("### 4.3 Marginal effect", "fig1_coefficients",
     "Nested model ladder. Point estimates with 95 per cent confidence intervals, "
     "clustered at district. The compound specification is highlighted.",
     True),
    ("### 4.4 Categorical decomposition", "fig2_marginal_effect",
     "Marginal effect of a one standard deviation moisture deficit on mean lean-season "
     "IPC phase, across growing-season heat conditions, 95 per cent confidence band. "
     "Where the pooled, whole-country estimate is not itself the strongest available "
     "(Section 4.7), shown alongside the same estimate within the stratum where the "
     "compound interaction is largest and individually significant, on a shared axis.",
     True),
    ("### 4.4 Categorical decomposition", "fig7_binned_interaction",
     "Nonparametric view of the same interaction. Both axes are two-way demeaned, so "
     "the panels show only the variation the fixed effects estimator uses. Points are "
     "binned means with 95 per cent confidence intervals.",
     False),
    ("### 4.5 Out-of-sample prediction", "fig3_categories",
     "Mean lean-season IPC phase by drought and heat co-occurrence, with 95 per cent "
     "confidence intervals. Categories are ordered by expected severity.",
     True),
    ("### 4.5 Out-of-sample prediction", "map2_exposure_outcome",
     "Geographic pattern. (a) Share of seasons classified as compound drought-heat. "
     "(b) Mean lean-season IPC phase. Hatched units carry no outcome data.",
     True),
    ("### 4.5 Out-of-sample prediction", "map3_hazard_frequency",
     "Component hazards mapped separately. (a) Share of seasons in drought alone. "
     "(b) Share of seasons in heat stress alone. Compare against [[FIG:map2_exposure_outcome]]a, "
     "compound exposure: a district can be prone to one hazard, both, or neither.",
     False),
    ("### 4.5 Out-of-sample prediction", "map_phase_periods",
     "Food insecurity by period. Small multiples on one shared colour scale, so the "
     "geographic pattern can be compared across periods directly rather than "
     "recalled from separate figures.",
     False),
    ("### 4.5 Out-of-sample prediction", "map_compound_periods",
     "Compound drought-heat exposure by period, same period bins and colour-scale "
     "convention as the food-insecurity-by-period figure.",
     False),
    ("### 4.5 Out-of-sample prediction", "fig8_spacetime",
     "The full record at once: every district-season's mean lean-season IPC phase, "
     "with compound drought-heat seasons marked ×. Districts are ordered by "
     "long-run rainfall climatology, driest at top.",
     False),
    ("### 4.6 Robustness", "fig4_cv",
     "Leave-one-season-out root mean squared error by specification. Lower is better. "
     "The held-out season is predicted from district intercepts and slopes only.",
     True),
    ("### 4.8 Spatial correlation", "map5_strata",
     "Agro-ecological stratification used in the heterogeneity analysis above.",
     True),
    ("### 4.8 Spatial correlation", "fig5_heterogeneity",
     "Compound interaction estimated separately within each agro-ecological stratum, "
     "with 95 per cent confidence intervals and sample sizes.",
     True),
    ("### 4.9 Rainfall product", "fig6_conley",
     "Compound interaction with Conley spatial HAC confidence intervals across spatial "
     "correlation cutoffs. The point estimate is invariant; only inference moves.",
     False),
    ("### 4.9 Rainfall product", "fig_placebo",
     "Randomization inference on the compound interaction: the null distribution "
     "from 500 within-district permutations of the season-year assignment of "
     "exposure (outcome held fixed), against the true estimate. A check with no "
     "asymptotic-normality assumption, independent of the clustered and Conley "
     "standard errors reported elsewhere.",
     False),
    ("## 5. Discussion", "map_scenario_delta",
     "Projected change in mean lean-season IPC phase by district under each CMIP6 "
     "scenario, 2050s against the 1995-2014 historical reference, delta method "
     "(Section 3.8). Diverging scale: some districts may project cooler or wetter.",
     True),
    ("## 5. Discussion", "fig9_scenario_burden",
     "Projected food-insecurity burden by the 2050s. (a) Population-weighted mean "
     "change in lean-season IPC phase. (b) Share of the reported population living "
     "in districts projected to cross a 0.5-phase-point increase. Population held "
     "at current levels in every scenario (a climate-only projection, not "
     "combined with demographic change).",
     False),
]

TABLES = [
    ("### 4.3 Marginal effect", "coefficients",
     "Nested model ladder: coefficients, standard errors clustered at district.",
     ["model", "term", "coef", "se", "p"]),
    ("### 4.6 Robustness", "cv",
     "Leave-one-season-out predictive performance.", None),
    ("### 4.7 Heterogeneity", "robustness",
     "Robustness of the compound interaction across specifications.", None),
    ("### 4.8 Spatial correlation", "heterogeneity",
     "Compound specification estimated within agro-ecological strata.", None),
    ("### 4.9 Rainfall product", "conley_sensitivity",
     "Conley spatial HAC standard errors across distance cutoffs.", None),
    ("### 4.10 Climate-change scenario", "rainfall_sensitivity",
     "Model ladder re-estimated on an independent rainfall product.", None),
    ("## 5. Discussion", "scenario_burden",
     "Projected population-weighted phase change and elevated-risk "
     "population share under each climate scenario, 2050s.", None),
]


def md_table(df: pd.DataFrame, cols=None, max_rows=40) -> str:
    if df.empty:
        return "_(not produced in this run)_"
    d = df[cols] if cols and set(cols).issubset(df.columns) else df
    d = d.head(max_rows).copy()
    for c in d.columns:
        if not pd.api.types.is_float_dtype(d[c]):
            continue
        if str(c).lower().startswith("p") and str(c).lower() in ("p", "p_interaction"):
            d[c] = d[c].map(lambda v: "" if pd.isna(v)
                            else ("< 0.001" if v < 0.001 else f"{v:.3f}"))
        else:
            d[c] = d[c].map(lambda v: "" if pd.isna(v) else f"{v:.4f}")
    d.columns = [str(c).replace("_", " ") for c in d.columns]
    head = "| " + " | ".join(d.columns) + " |"
    rule = "| " + " | ".join("---" for _ in d.columns) + " |"
    body = "\n".join("| " + " | ".join(str(x) for x in row) + " |"
                     for row in d.itertuples(index=False))
    return "\n".join([head, rule, body])


def _find_pandoc() -> str:
    """Locate the pandoc executable.

    shutil.which("pandoc") fails right after a fresh install on Windows: the
    installer writes the user PATH to the registry, but a shell already open
    at that point (this one, reliably) keeps its own stale copy until it is
    restarted. Fall back to the default per-user install location before
    concluding pandoc is actually missing.
    """
    found = shutil.which("pandoc")
    if found:
        return found
    fallback = Path.home() / "AppData" / "Local" / "Pandoc" / "pandoc.exe"
    if fallback.exists():
        return str(fallback)
    raise FileNotFoundError(
        "pandoc not found on PATH or at the default per-user install location "
        f"({fallback}). Install it from https://pandoc.org/installing.html, or "
        "if you just installed it, open a new terminal so it picks up PATH.")


def _set_cell_border(cell, **edges):
    """Set individual border edges on one table cell (python-docx has no
    higher-level API for this). edges: e.g. top={"sz": 12, "color": "000000"}."""
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement
    tcPr = cell._tc.get_or_add_tcPr()
    borders = tcPr.find(qn("w:tcBorders"))
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tcPr.append(borders)
    for edge, spec in edges.items():
        tag = f"w:{edge}"
        el = borders.find(qn(tag))
        if el is None:
            el = OxmlElement(tag)
            borders.append(el)
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), str(spec["sz"]))
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), spec.get("color", "000000"))


def _style_table_academic(table):
    """Three-line ("booktabs") table styling: a rule above the header, a
    rule below it, and a rule at the bottom -- nothing else. Pandoc's own
    docx output carries no borders at all (the base Word "Normal Table"
    style is borderless) and no header emphasis, so a multi-row results
    table currently reads as plain unaligned text with no header line and
    no visual separation between rows -- worse than a spreadsheet grid, not
    merely different from one. This is the standard journal convention
    (Elsevier, APA, LaTeX booktabs): horizontal rules only, never vertical
    lines, never rules between ordinary data rows.
    """
    thick, thin = 12, 6  # eighths of a point: 1.5pt and 0.75pt
    header = table.rows[0]
    for cell in header.cells:
        _set_cell_border(cell, top={"sz": thick}, bottom={"sz": thin})
        for p in cell.paragraphs:
            for run in p.runs:
                run.bold = True
    for cell in table.rows[-1].cells:
        _set_cell_border(cell, bottom={"sz": thick})


def ordered_entries(text: str, entries, directory: Path, suffix: str, tag: str):
    """Use the same verified document order for captions and packaged assets."""
    lines = text.splitlines()
    def position(entry):
        matches = [i for i, line in enumerate(lines) if line.startswith(entry[0])]
        if len(matches) != 1:
            raise ValueError(f"Expected one anchor for {entry[1]}: {entry[0]!r}; found {len(matches)}")
        return matches[0]
    return sorted((e for e in entries if (directory / f"{e[1]}_{tag}.{suffix}").exists()),
                  key=position)


def _split_main_supp(fig_entries):
    """FIGURES entries carry a trailing `main` flag (True = manuscript body,
    False = Supplementary Material). Splitting preserves each sub-list's
    document order -- the input is already anchor-order -- so numbering
    within each tier is anchor-order, not array-order."""
    return ([e for e in fig_entries if e[-1]],
            [e for e in fig_entries if not e[-1]])


def _build_numbers(fig_entries, tab_entries) -> dict:
    """Figure numbering splits into a main-text run (Figure 1, 2, ...) and a
    Supplementary Material run (Figure S1, S2, ...), each numbered by its
    own position in document order. Tables are one run, always main-text.
    Building this over ALL entries (not just the ones landing in whichever
    document is currently being built) is what lets a main-text caption
    cite a supplementary figure, and vice versa, and have both resolve.
    """
    fig_main, fig_supp = _split_main_supp(fig_entries)
    return {
        "FIG": {**{e[1]: f"Figure {n}" for n, e in enumerate(fig_main, 1)},
                **{e[1]: f"Figure S{n}" for n, e in enumerate(fig_supp, 1)}},
        "TAB": {e[1]: f"Table {n}" for n, e in enumerate(tab_entries, 1)},
    }


def resolve_references(text: str, fig_entries, tab_entries) -> str:
    numbers = _build_numbers(fig_entries, tab_entries)
    def replace(match):
        kind, stem = match.groups()
        if stem not in numbers[kind]:
            raise ValueError(f"Referenced {kind} asset is missing: {stem}")
        return numbers[kind][stem]
    return re.sub(r"\[\[(FIG|TAB):([a-z0-9_]+)\]\]", replace, text)


def build(tag: str) -> Path:
    src = C.ROOT / "manuscript" / f"manuscript_filled_{tag}.md"
    if not src.exists():
        raise FileNotFoundError(f"{src} not found. Run fill_manuscript.py first.")
    text = src.read_text(encoding="utf-8")
    blocks: dict[str, list[str]] = {}
    fig_entries = ordered_entries(text, FIGURES, C.FIG, "png", tag)
    tab_entries = ordered_entries(text, TABLES, C.TAB, "csv", tag)
    numbers = _build_numbers(fig_entries, tab_entries)
    fig_main, _fig_supp = _split_main_supp(fig_entries)

    text = resolve_references(text, fig_entries, tab_entries)
    lines = text.split("\n")
    for anchor, stem, caption, _main in fig_main:
        png = C.FIG / f"{stem}_{tag}.png"
        caption = resolve_references(caption, fig_entries, tab_entries)
        blocks.setdefault(anchor, []).append(
            f"\n![]({png.as_posix()}){{width=6.2in}}\n\n*{numbers['FIG'][stem]}. {caption}*\n")
    for anchor, name, caption, cols in tab_entries:
        f = C.TAB / f"{name}_{tag}.csv"
        caption = resolve_references(caption, fig_entries, tab_entries)
        blocks.setdefault(anchor, []).append(
            f"\n*{numbers['TAB'][name]}. {caption}*\n\n{md_table(pd.read_csv(f), cols)}\n")
    if _fig_supp:
        blocks.setdefault("## References", []).insert(0,
            "\nSee the Supplementary Material for further figures and "
            "robustness checks.\n")

    # Insert each block just BEFORE the anchor heading that follows its section.
    out = []
    for line in lines:
        for anchor, chunks in blocks.items():
            if line.startswith(anchor):
                out.extend(chunks)
                blocks[anchor] = []
        out.append(line)
    # anything unplaced goes at the end
    for chunks in blocks.values():
        out.extend(chunks)

    staged = C.ROOT / "manuscript" / f"_staged_{tag}.md"
    staged.write_text("\n".join(out), encoding="utf-8")

    dst = C.ROOT / "manuscript" / f"manuscript_{tag}.docx"
    cmd = [
        _find_pandoc(), str(staged), "-o", str(dst),
        "--from", "markdown+pipe_tables+tex_math_dollars",
        "--to", "docx",
        # No --toc: a table of contents is a generated-document tell a journal
        # submission does not carry. Section numbering in the headings does
        # the same navigational job.
        "--resource-path", str(C.ROOT),
    ]
    # Fonts/heading styles only (Times New Roman body, bold Times headings),
    # lifted from the user's own GHAP/CAMS manuscript's styles.xml/theme1.xml
    # onto pandoc's blank scaffold -- no content or embedded figures from that
    # paper, just its Word template.
    ref_doc = C.ROOT / "manuscript" / "template" / "reference_style.docx"
    if ref_doc.exists():
        cmd += ["--reference-doc", str(ref_doc)]
    subprocess.run(cmd, check=True)
    _polish_docx(dst)
    staged.unlink(missing_ok=True)
    return dst


def _set_run_font(style_or_run, name):
    """Set a font by name on every script Word tracks separately (Latin,
    East Asian, complex-script). Setting only python-docx's font.name leaves
    w:eastAsia untouched, which some Word versions fall back to for glyphs
    it decides aren't plain Latin, undoing the font change unpredictably."""
    from docx.oxml.ns import qn
    style_or_run.font.name = name
    rPr = style_or_run.element.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = rPr.makeelement(qn("w:rFonts"), {})
        rPr.append(rFonts)
    for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rFonts.set(qn(attr), name)


def _polish_docx(dst: Path, manuscript_conventions: bool = True):
    """Page setup, font, justification, image sizing, keep-together, and
    table styling -- everything pandoc's own docx output does not do on its
    own. Shared by build(), build_supplementary() and the cover letter build
    so every document gets identical base formatting.

    manuscript_conventions adds continuous line numbers and double spacing:
    both are peer-review-submission conventions (a reviewer cites a line
    number; a line needs room for a handwritten comment) that make sense for
    the manuscript and its supplement but not for a one-page cover letter,
    which nobody reviews line-by-line.
    """
    from docx import Document
    from docx.shared import Inches, Mm, Pt
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement
    doc = Document(dst)
    # Arial throughout, overriding the reference template's Times New Roman:
    # set at the style level first (so it's the default for any paragraph
    # using that style, table cells included, since cells use "Normal" too
    # but are not in doc.paragraphs), then forced at the run level as well,
    # because pandoc sometimes bakes an explicit run-level font that would
    # otherwise shadow the style default. Body text is additionally forced
    # to exactly 12pt; headings keep their own template sizes.
    for style_name in ("Normal", "Body Text", "Heading 1", "Heading 2", "Heading 3", "Title"):
        if style_name in doc.styles:
            _set_run_font(doc.styles[style_name], "Arial")
    for style_name in ("Normal", "Body Text"):
        if style_name in doc.styles:
            doc.styles[style_name].font.size = Pt(12)
    all_paragraphs = list(doc.paragraphs)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                all_paragraphs.extend(cell.paragraphs)
    for paragraph in all_paragraphs:
        for run in paragraph.runs:
            _set_run_font(run, "Arial")
    # Justified body text (headings stay left-aligned, the normal exception).
    for paragraph in doc.paragraphs:
        if paragraph.style.name in ("Normal", "Body Text"):
            paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    # APA hanging indent on the reference list: markdown has no such concept,
    # so every entry pandoc produces starts flush left. Applied to every
    # body paragraph from the "References" heading onward; harmless no-op
    # in a document with no such heading (the supplement, the cover letter).
    in_references = False
    for paragraph in doc.paragraphs:
        if paragraph.style.name.startswith("Heading") and paragraph.text.strip() == "References":
            in_references = True
            continue
        if in_references and paragraph.style.name in ("Normal", "Body Text") and paragraph.text.strip():
            paragraph.paragraph_format.left_indent = Inches(0.5)
            paragraph.paragraph_format.first_line_indent = Inches(-0.5)
    for section in doc.sections:
        section.page_width = Mm(210)
        section.page_height = Mm(297)
        section.left_margin = section.right_margin = Inches(1)
        section.top_margin = section.bottom_margin = Inches(1)
        if manuscript_conventions:
            # Continuous line numbers: a peer-review submission convention
            # (reviewers cite "line 214", not a paragraph) that a markdown
            # export has no way to produce on its own.
            sectPr = section._sectPr
            lnNumType = OxmlElement("w:lnNumType")
            lnNumType.set(qn("w:countBy"), "1")
            lnNumType.set(qn("w:restart"), "continuous")
            sectPr.append(lnNumType)
    if manuscript_conventions:
        # Double-spaced body text, the standard review-submission convention
        # (headings stay single-spaced, the normal exception). Table cells
        # use the "Normal" style too, so this sets spacing per body
        # paragraph (doc.paragraphs excludes table-cell paragraphs) rather
        # than on the style itself -- doing it at the style level would
        # double-space every table row as a side effect.
        for paragraph in doc.paragraphs:
            if paragraph.style.name in ("Normal", "Body Text"):
                paragraph.paragraph_format.line_spacing = 2.0
    # Tall heatmaps must fit on one page with their caption. A width-only
    # constraint can produce an 11-inch image and endless Word repagination.
    for shape in doc.inline_shapes:
        scale = min(1.0, Inches(6.2) / shape.width, Inches(8.0) / shape.height)
        shape.width = int(shape.width * scale)
        shape.height = int(shape.height * scale)
    for paragraph in doc.paragraphs:
        if paragraph._p.xpath(".//w:drawing"):
            paragraph.paragraph_format.keep_with_next = True
        if re.match(r"(?:Figure|Table) S?\d+\.", paragraph.text):
            paragraph.paragraph_format.keep_together = True
            paragraph.paragraph_format.keep_with_next = paragraph.text.startswith("Table ")
    for table in doc.tables:
        _style_table_academic(table)
    doc.save(dst)


def build_supplementary(tag: str) -> Path | None:
    """A separate Supplementary Material document holding every figure not
    promoted to the manuscript body, numbered Figure S1, S2, ... in the same
    document order the main text's cross-references to them assume.
    """
    src = C.ROOT / "manuscript" / f"manuscript_filled_{tag}.md"
    text = src.read_text(encoding="utf-8")
    fig_entries = ordered_entries(text, FIGURES, C.FIG, "png", tag)
    tab_entries = ordered_entries(text, TABLES, C.TAB, "csv", tag)
    numbers = _build_numbers(fig_entries, tab_entries)
    _fig_main, fig_supp = _split_main_supp(fig_entries)
    if not fig_supp:
        return None

    out = ["# Supplementary Material", "",
           "Compound drought and heat exposure and lean-season food insecurity "
           "in Zimbabwe, 2009 to 2025", "", ""]
    for anchor, stem, caption, _main in fig_supp:
        png = C.FIG / f"{stem}_{tag}.png"
        caption = resolve_references(caption, fig_entries, tab_entries)
        out += [f"![]({png.as_posix()}){{width=6.2in}}", "",
                f"*{numbers['FIG'][stem]}. {caption}*", ""]

    staged = C.ROOT / "manuscript" / f"_staged_supplementary_{tag}.md"
    staged.write_text("\n".join(out), encoding="utf-8")
    dst = C.ROOT / "manuscript" / f"manuscript_supplementary_{tag}.docx"
    cmd = [_find_pandoc(), str(staged), "-o", str(dst),
           "--from", "markdown+pipe_tables+tex_math_dollars", "--to", "docx",
           "--resource-path", str(C.ROOT)]
    ref_doc = C.ROOT / "manuscript" / "template" / "reference_style.docx"
    if ref_doc.exists():
        cmd += ["--reference-doc", str(ref_doc)]
    subprocess.run(cmd, check=True)
    _polish_docx(dst)
    staged.unlink(missing_ok=True)
    return dst


def verify(docx: Path):
    """Render and report page count, so a broken export is visible."""
    if not shutil.which("soffice"):
        return None
    outdir = docx.parent / "_render"
    outdir.mkdir(exist_ok=True)
    subprocess.run(["soffice", "--headless", "--convert-to", "pdf",
                    "--outdir", str(outdir), str(docx)],
                   check=False, capture_output=True, timeout=240)
    pdf = outdir / (docx.stem + ".pdf")
    return pdf if pdf.exists() else None


if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("--synthetic", action="store_true")
    a = ap.parse_args()
    tag = "synthetic" if a.synthetic else "real"
    out = build(tag)
    print(f"Wrote {out}  ({out.stat().st_size/1024:.0f} KB)")
    pdf = verify(out)
    if pdf:
        print(f"Rendered check: {pdf}")
    supp = build_supplementary(tag)
    if supp:
        print(f"Wrote {supp}  ({supp.stat().st_size/1024:.0f} KB)")
