"""Check figure order, preceding callouts, image identities and package hashes."""
import hashlib
import json
import re
from pathlib import Path
from docx import Document
from export_docx import C, FIGURES, TABLES, ordered_entries, _split_main_supp


def _verify_document(docx_path, fig_entries, tab_entries=(), require_callout=True):
    """fig_entries/tab_entries: the entries THIS document contains, already
    in the exact order they are captioned in (Figure 1, 2, ... or S1, S2, ...).

    require_callout=False for the Supplementary Material: by convention it
    presents each item without re-narrating it, since the callout that sent
    the reader there lives in the main text instead.
    """
    doc = Document(docx_path)
    paragraphs = doc.paragraphs
    seen = {"Figure": [], "Table": []}
    for index, paragraph in enumerate(paragraphs):
        match = re.match(r"^(Figure|Table) (S?\d+)\.", paragraph.text)
        if match:
            kind, label = match[1], match[2]
            number = int(label.lstrip("S"))
            if require_callout:
                prior = "\n".join(p.text for p in paragraphs[:index]
                                  if not re.match(r"^(Figure|Table) S?\d+\.", p.text))
                assert re.search(rf"\b{kind} {re.escape(label)}\b", prior), \
                    f"{docx_path.name}: no preceding callout for {kind} {label}"
            seen[kind].append(number)
            if kind == "Figure":
                image_paragraph = paragraphs[index - 1]
                blips = image_paragraph._p.xpath(".//a:blip")
                assert len(blips) == 1, \
                    f"{docx_path.name}: {kind} {label} is detached from its image"
                relationship = blips[0].get(
                    "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed")
                expected = (C.FIG / f"{fig_entries[number-1][1]}_real.png").read_bytes()
                assert doc.part.related_parts[relationship].blob == expected, \
                    f"{docx_path.name}: {kind} {label} is stale"
    assert seen["Figure"] == list(range(1, len(fig_entries) + 1)), (docx_path.name, seen)
    assert seen["Table"] == list(range(1, len(tab_entries) + 1)), (docx_path.name, seen)
    assert len(doc.inline_shapes) == len(fig_entries), docx_path.name
    assert all(s.height <= 8.0 * 914400 for s in doc.inline_shapes), \
        f"{docx_path.name}: an image exceeds its page budget"
    assert len(doc.tables) == len(tab_entries), docx_path.name
    assert not re.search(r"\[\[(?:[A-Z0-9_]+|(?:FIG|TAB):[a-z0-9_]+)\]\]",
                         "\n".join(p.text for p in paragraphs)), \
        f"{docx_path.name}: unresolved marker left in the document"


def verify(package=False):
    source = (C.ROOT / "manuscript/manuscript_filled_real.md").read_text(encoding="utf-8")
    figures = ordered_entries(source, FIGURES, C.FIG, "png", "real")
    tables = ordered_entries(source, TABLES, C.TAB, "csv", "real")
    fig_main, fig_supp = _split_main_supp(figures)

    _verify_document(C.ROOT / "manuscript/manuscript_real.docx", fig_main, tables)
    if fig_supp:
        _verify_document(C.ROOT / "manuscript/manuscript_supplementary_real.docx",
                         fig_supp, require_callout=False)

    if package:
        dest = C.ROOT / "submission"
        manifest = json.loads((dest / "asset_manifest.json").read_text(encoding="utf-8"))
        assert len(manifest) == len(figures) + len(tables)
        for entry in manifest:
            assert hashlib.sha256((dest / entry["file"]).read_bytes()).hexdigest() == entry["sha256"]
        assert (dest / "manuscript.docx").read_bytes() == \
            (C.ROOT / "manuscript/manuscript_real.docx").read_bytes()
        if fig_supp:
            assert (dest / "manuscript_supplementary.docx").read_bytes() == \
                (C.ROOT / "manuscript/manuscript_supplementary_real.docx").read_bytes()
    print(f"PASS: {len(fig_main)} main + {len(fig_supp)} supplementary figures, "
          f"{len(tables)} tables, sequential captions, preceding callouts, current images"
          + (", package checksums" if package else ""))


if __name__ == "__main__":
    import sys
    verify(package="--package" in sys.argv)
